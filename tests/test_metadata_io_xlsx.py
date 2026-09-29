"""`data/metadata_profile/io_xlsx.py` 的工作簿读写、备份与原子落盘测试。

工作簿用 openpyxl 手工搭建（只依赖 `layout`），因此不依赖生成器脚本，也不依赖 markdown 解析。
"""
from __future__ import annotations

from datetime import date, datetime
from pathlib import Path

import openpyxl
import pytest
from openpyxl import Workbook

from data.metadata_profile.io_xlsx import (
    CHECK_HEADERS,
    FIRST_DATA_ROW,
    backup_file,
    read_sheet,
    read_sheets,
    write_check_report,
    write_records,
)
from data.metadata_profile.layout import (
    SHEET_CHECK,
    SHEET_README,
    SHEET_SCHEMA,
    SHEET_VOCAB,
)


def _build_workbook(path: Path, layout) -> Path:
    """按布局搭一个空模板（仅表头，供读取/写入测试）。"""
    workbook = Workbook()
    workbook.remove(workbook.active)
    for sheet_spec in layout.sheets:
        worksheet = workbook.create_sheet(title=sheet_spec.name)
        for column, key in enumerate(sheet_spec.all_header_keys, start=1):
            worksheet.cell(row=1, column=column, value=key)
        for column, label in enumerate(sheet_spec.all_header_labels, start=1):
            worksheet.cell(row=2, column=column, value=label)
        if sheet_spec.name == SHEET_CHECK:
            for column, header in enumerate(CHECK_HEADERS, start=1):
                worksheet.cell(row=1, column=column, value=header)
    workbook.save(path)
    return path


@pytest.fixture()
def template(tmp_path, synthetic_layout):
    """An empty, layout-faithful workbook on disk."""
    return _build_workbook(tmp_path / "entry.xlsx", synthetic_layout)


def test_read_sheets_maps_headers_to_machine_keys(template, synthetic_layout):
    """读取以布局为准，按表头第 1 行的机器键名建键。"""
    result = read_sheets(synthetic_layout, template)
    assert set(result) == {s.name for s in synthetic_layout.entry_sheets}
    site = result["1_Site"]
    assert site.table == "sites"
    assert site.records == ()  # 空模板没有数据行
    analysis = result["5_Analysis"]
    assert analysis.table == "analyses"


def test_read_sheet_skips_blank_rows_and_normalizes_values(
    template, synthetic_layout
):
    """空行跳过；日期归一为 YYYY-MM-DD；空白字符串归一为 None。"""
    workbook = openpyxl.load_workbook(template)
    worksheet = workbook["1_Site"]
    row = FIRST_DATA_ROW
    worksheet.cell(row=row, column=1, value="  SI-0001  ")  # 去空白
    worksheet.cell(row=row, column=2, value="   ")          # 空白 -> None
    # 第 row+1 行整行留空 -> 跳过
    worksheet.cell(row=row + 2, column=1, value="SI-0002")
    workbook.save(template)

    sheet_spec = synthetic_layout.sheet("1_Site")
    reread = openpyxl.load_workbook(template)["1_Site"]
    result = read_sheet(reread, sheet_spec)

    assert [r["SI0 terralid_site_id"] for r in result.records] == [
        "SI-0001",
        "SI-0002",
    ]
    assert result.records[0]["SI1 site_name"] is None
    assert result.records[0]["__row__"] == row


def test_read_sheet_converts_excel_dates_to_iso(template, synthetic_layout):
    """Excel 日期单元格转成档案要求的 YYYY-MM-DD 字符串。"""
    workbook = openpyxl.load_workbook(template)
    worksheet = workbook["5_Analysis"]
    worksheet.cell(row=FIRST_DATA_ROW, column=1, value="A-1")
    # A12 分析日期所在列
    keys = list(synthetic_layout.sheet("5_Analysis").all_header_keys)
    date_col = keys.index("A12 analysis_lia_date") + 1
    worksheet.cell(row=FIRST_DATA_ROW, column=date_col, value=datetime(2024, 2, 24))
    workbook.save(template)

    result = read_sheet(
        openpyxl.load_workbook(template)["5_Analysis"],
        synthetic_layout.sheet("5_Analysis"),
    )
    assert result.records[0]["A12 analysis_lia_date"] == "2024-02-24"


def test_read_sheet_ignores_extra_columns(template, synthetic_layout):
    """用户在布局之外自己加的列必须被忽略，不污染记录。"""
    workbook = openpyxl.load_workbook(template)
    worksheet = workbook["1_Site"]
    extra = len(synthetic_layout.sheet("1_Site").all_header_keys) + 1
    worksheet.cell(row=1, column=extra, value="my own note")
    worksheet.cell(row=FIRST_DATA_ROW, column=1, value="SI-1")
    worksheet.cell(row=FIRST_DATA_ROW, column=extra, value="忽略我")
    workbook.save(template)

    result = read_sheet(
        openpyxl.load_workbook(template)["1_Site"],
        synthetic_layout.sheet("1_Site"),
    )
    assert "my own note" not in result.records[0]
    assert result.records[0]["SI0 terralid_site_id"] == "SI-1"


def test_read_sheets_rejects_missing_file(synthetic_layout, tmp_path):
    """文件不存在时抛 FileNotFoundError（而不是静默返回空）。"""
    with pytest.raises(FileNotFoundError):
        read_sheets(synthetic_layout, tmp_path / "nope.xlsx")


def test_read_sheets_rejects_non_xlsx(synthetic_layout, tmp_path):
    """非 xlsx 内容抛 ValueError 并记日志。"""
    bogus = tmp_path / "bogus.xlsx"
    bogus.write_text("not a workbook", encoding="utf-8")
    with pytest.raises(ValueError):
        read_sheets(synthetic_layout, bogus)


def test_backup_file_creates_timestamped_copy(tmp_path):
    """备份文件名带时间戳，且内容一致；源文件不存在时返回 None。"""
    source = tmp_path / "data.xlsx"
    source.write_bytes(b"original")
    backup = backup_file(source, timestamp="20260101-120000")
    assert backup is not None
    assert backup.name == "data.bak-20260101-120000.xlsx"
    assert backup.read_bytes() == b"original"
    assert backup_file(tmp_path / "missing.xlsx") is None


def test_write_check_report_writes_rows_and_backs_up(template, synthetic_layout):
    """校验报告写入 99_CHECK，并在覆盖前留下备份。"""
    issues = [
        {
            "severity": "error",
            "sheet": "1_Site",
            "record_id": "SI-1",
            "column": "SI1 site_name",
            "message": "Required field is empty",
            "value": None,
        },
        {
            "severity": "warning",
            "sheet": "5_Analysis",
            "record_id": "A-1",
            "column": "A12 analysis_lia_date",
            "message": "Date is not YYYY-MM-DD",
            "value": "24/02/2024",
        },
    ]
    target, backup = write_check_report(synthetic_layout, template, issues)
    assert target == template
    assert backup is not None and backup.exists()

    worksheet = openpyxl.load_workbook(template)[SHEET_CHECK]
    assert [c.value for c in worksheet[1]][: len(CHECK_HEADERS)] == list(CHECK_HEADERS)
    assert worksheet.cell(row=2, column=1).value == "error"
    assert worksheet.cell(row=3, column=4).value == "A12 analysis_lia_date"


def test_write_check_report_is_idempotent(template, synthetic_layout):
    """重复写报告不累积旧行。"""
    many = [
        {
            "severity": "error",
            "sheet": "1_Site",
            "record_id": f"SI-{i}",
            "column": "SI1 site_name",
            "message": "Required field is empty",
            "value": None,
        }
        for i in range(5)
    ]
    write_check_report(synthetic_layout, template, many)
    write_check_report(synthetic_layout, template, many[:1])
    worksheet = openpyxl.load_workbook(template)[SHEET_CHECK]
    rows = [
        r for r in worksheet.iter_rows(min_row=2, values_only=True) if any(r)
    ]
    assert len(rows) == 1


def test_write_records_round_trips_provider_values(template, synthetic_layout):
    """写入记录后能原样读回（键为机器键名）。"""
    records = [
        {
            "SI0 terralid_site_id": "SI-0001",
            "SI1 site_name": "Anyang",
            "SI5 site_geolocation": 36.1,
        }
    ]
    _, backup = write_records(synthetic_layout, template, {"1_Site": records})
    assert backup is not None

    read_back = read_sheets(synthetic_layout, template, sheets=["1_Site"])
    assert read_back["1_Site"].records[0]["SI1 site_name"] == "Anyang"


def test_write_records_refuses_generated_sheets(template, synthetic_layout):
    """生成物表（0_SCHEMA / README / 99_CHECK）不得被写入记录。"""
    for name in (SHEET_SCHEMA, SHEET_README, SHEET_VOCAB, SHEET_CHECK):
        with pytest.raises(ValueError):
            write_records(synthetic_layout, template, {name: [{"x": 1}]})


def test_write_records_rejects_unknown_sheet(template, synthetic_layout):
    """布局里不存在的表名必须报错，避免写错表。"""
    with pytest.raises(ValueError):
        write_records(synthetic_layout, template, {"99_Nonexistent": []})


def test_write_records_replaces_previous_rows(template, synthetic_layout):
    """再次写入时先清空旧数据行，不残留。"""
    write_records(
        synthetic_layout,
        template,
        {"1_Site": [{"SI0 terralid_site_id": f"SI-{i}"} for i in range(4)]},
    )
    write_records(
        synthetic_layout, template, {"1_Site": [{"SI0 terralid_site_id": "SI-9"}]}
    )
    result = read_sheets(synthetic_layout, template, sheets=["1_Site"])
    assert [r["SI0 terralid_site_id"] for r in result["1_Site"].records] == ["SI-9"]


def test_date_value_with_date_object(tmp_path, synthetic_layout):
    """``date`` 对象同样归一为 ISO 字符串。"""
    template = _build_workbook(tmp_path / "d.xlsx", synthetic_layout)
    workbook = openpyxl.load_workbook(template)
    worksheet = workbook["5_Analysis"]
    keys = list(synthetic_layout.sheet("5_Analysis").all_header_keys)
    worksheet.cell(row=FIRST_DATA_ROW, column=1, value="A-1")
    worksheet.cell(
        row=FIRST_DATA_ROW,
        column=keys.index("A12 analysis_lia_date") + 1,
        value=date(2024, 2, 24),
    )
    workbook.save(template)
    result = read_sheet(
        openpyxl.load_workbook(template)["5_Analysis"],
        synthetic_layout.sheet("5_Analysis"),
    )
    assert result.records[0]["A12 analysis_lia_date"] == "2024-02-24"
