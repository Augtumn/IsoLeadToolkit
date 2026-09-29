"""`scripts/build_entry_workbook.py` 的录入工作簿生成测试。

主体用共享夹具 `synthetic_profile` / `synthetic_layout`（`tests/conftest.py`，由 `spec.py`
手工构造、不经 markdown 解析），保证生成器本身可独立验证；文件末尾另有一组真实档案
（`profile.load_profile()`）的端到端断言，覆盖只有真档案才有的列（GEOL_TYPE / UNIT_AGE /
YES_NO_UNCLEAR、双宿主的 `ratio_host`）。

批注尺寸只在**内存态**工作簿上可读（openpyxl 往返不保留批注框尺寸），因此这类断言走
`_build_workbook()`；其余断言都走公开 API `workbook_bytes()`。
"""
from __future__ import annotations

import json
from io import BytesIO
from pathlib import Path

import pytest
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

import scripts.build_entry_workbook as build_entry_workbook
from data.metadata_profile.io_xlsx import CHECK_HEADERS, FIRST_DATA_ROW
from data.metadata_profile.layout import (
    OWNER_SHEET_COLUMN,
    RATIO_HOST_COLUMN,
    SHEET_CHECK,
    SHEET_README,
    SHEET_SCHEMA,
    SHEET_VOCAB,
    SheetKind,
    build_layout,
    owner_sheet_options,
)
from data.metadata_profile.spec import (
    LIA_RATIO_NAMES,
    RATIO_SOURCES,
    SIGMA_LEVELS,
    Block,
    Profile,
    ValueKind,
)
from scripts.build_entry_workbook import (
    ENTRY_ROW_COUNT,
    SCHEMA_HEADERS,
    _build_workbook,
    build_workbook,
    main,
    schema_rows,
    vocab_layout,
    workbook_bytes,
)

LAST_DATA_ROW = FIRST_DATA_ROW + ENTRY_ROW_COUNT - 1


def _workbook(profile: Profile):
    """生成并读回工作簿（公开 API 往返）。"""
    return load_workbook(BytesIO(workbook_bytes(profile)))


def _column_index(sheet_spec, header_key: str) -> int:
    """表头键 → 列号（1 起）。"""
    return list(sheet_spec.all_header_keys).index(header_key) + 1


def _validation_for(worksheet: Worksheet, column_index: int):
    """取覆盖某列的数据验证（list 型下拉）；没有则返回 ``None``。"""
    for validation in worksheet.data_validations.dataValidation:
        for cell_range in validation.sqref.ranges:
            if cell_range.min_col <= column_index <= cell_range.max_col:
                return validation
    return None


def _conditional_rules_for(worksheet: Worksheet, column_index: int) -> list[object]:
    """取覆盖某列的条件格式规则。"""
    rules: list[object] = []
    for formatting in worksheet.conditional_formatting:
        for cell_range in formatting.sqref.ranges:
            if cell_range.min_col <= column_index <= cell_range.max_col:
                rules.extend(formatting.rules)
    return rules


def _expected_schema_row_count(profile: Profile) -> int:
    """独立算出 `0_SCHEMA` 应有的行数（模块自有字段 + 内联块字段 × 宿主数）。"""
    total = 0
    for table in profile.tables:
        total += len({spec.oid for spec in table.fields if spec.block is None})
        hosts: dict[Block, set[str]] = {}
        oids: dict[Block, set[str]] = {}
        for spec in table.fields:
            if spec.block is None or not spec.block_owner:
                continue
            hosts.setdefault(spec.block, set()).add(spec.block_owner)
            oids.setdefault(spec.block, set()).add(spec.oid)
        for block, host_set in hosts.items():
            total += len(host_set) * len(oids[block])
    return total


# ── 工作簿骨架 ──────────────────────────────────────────────────────────────


def test_workbook_bytes_round_trips_through_openpyxl(synthetic_profile):
    """`workbook_bytes()` 的结果能被 openpyxl 从内存读回。"""
    data = workbook_bytes(synthetic_profile)
    assert data[:2] == b"PK"  # xlsx 即 zip
    workbook = load_workbook(BytesIO(data))
    assert workbook.sheetnames


def test_sheet_names_and_order_follow_layout(synthetic_profile, synthetic_layout):
    """工作表名与顺序 == `layout.build_layout(profile).sheet_names`。"""
    workbook = _workbook(synthetic_profile)
    assert workbook.sheetnames == list(synthetic_layout.sheet_names)
    assert workbook.sheetnames[0] == SHEET_README
    assert workbook.sheetnames[-1] == SHEET_CHECK


def test_entry_sheets_have_two_row_header_and_frozen_panes(synthetic_profile, synthetic_layout):
    """录入表：第 1 行机器键名、第 2 行中文标签 + 级别标记、冻结 A3。"""
    workbook = _workbook(synthetic_profile)
    for sheet_spec in synthetic_layout.entry_sheets:
        worksheet = workbook[sheet_spec.name]
        keys = [worksheet.cell(row=1, column=i).value for i in range(1, sheet_spec.column_count + 1)]
        labels = [worksheet.cell(row=2, column=i).value for i in range(1, sheet_spec.column_count + 1)]
        assert keys == list(sheet_spec.all_header_keys), sheet_spec.name
        assert labels == list(sheet_spec.all_header_labels), sheet_spec.name
        assert worksheet.freeze_panes == "A3", sheet_spec.name


def test_first_header_row_has_no_empty_cells(synthetic_profile, synthetic_layout):
    """每张表的第 1 行在自己的列宽内没有空单元格。"""
    workbook = _workbook(synthetic_profile)
    widths = {
        SHEET_SCHEMA: len(SCHEMA_HEADERS),
        SHEET_CHECK: len(CHECK_HEADERS),
        SHEET_VOCAB: workbook[SHEET_VOCAB].max_column,
        SHEET_README: 1,
    }
    for sheet_spec in synthetic_layout.entry_sheets:
        widths[sheet_spec.name] = sheet_spec.column_count
    for name, column_count in widths.items():
        worksheet = workbook[name]
        values = [worksheet.cell(row=1, column=i).value for i in range(1, column_count + 1)]
        assert all(value not in (None, "") for value in values), name


def test_column_widths_come_from_layout(synthetic_profile, synthetic_layout):
    """列宽取 `ColumnSpec.width`，辅助列落在 16–22。"""
    workbook = _workbook(synthetic_profile)
    sheet_spec = synthetic_layout.sheet_for_block(Block.LIA_RATIO)
    worksheet = workbook[sheet_spec.name]
    for index, name in enumerate(sheet_spec.helper_columns, start=1):
        width = worksheet.column_dimensions[get_column_letter(index)].width
        assert 16 <= width <= 22, name
    for index in range(len(sheet_spec.helper_columns) + 1, sheet_spec.column_count + 1):
        column = sheet_spec.columns[index - len(sheet_spec.helper_columns) - 1]
        assert worksheet.column_dimensions[get_column_letter(index)].width == column.width


# ── 数值格式 ────────────────────────────────────────────────────────────────


def test_ratio_and_uncertainty_number_formats_on_all_data_rows(
    synthetic_profile, synthetic_layout
):
    """B6.2 用 5 位小数、B6.5 用 8 位小数，且覆盖全部预设数据行。"""
    workbook = _workbook(synthetic_profile)
    sheet_spec = synthetic_layout.sheet_for_block(Block.LIA_RATIO)
    worksheet = workbook[sheet_spec.name]

    value_column = _column_index(sheet_spec, "B6.2 lia_ratio_value")
    uncertainty_column = _column_index(
        sheet_spec, "B6.5 lia_ratio_uncertainty_value_absolute"
    )
    assert sheet_spec.column("B6.2").number_format == "0.00000"
    assert sheet_spec.column("B6.5").number_format == "0.00000000"
    for row in (FIRST_DATA_ROW, 100, LAST_DATA_ROW):
        assert worksheet.cell(row=row, column=value_column).number_format == "0.00000"
        assert (
            worksheet.cell(row=row, column=uncertainty_column).number_format
            == "0.00000000"
        )
    name_column = _column_index(sheet_spec, "B6.1 lia_ratio_name")
    assert worksheet.cell(row=FIRST_DATA_ROW, column=name_column).number_format == "General"


def test_date_columns_use_iso_number_format(synthetic_profile, synthetic_layout):
    """日期列按档案要求用 `yyyy-mm-dd`。"""
    workbook = _workbook(synthetic_profile)
    analysis = synthetic_layout.sheet_for_table("analyses")
    worksheet = workbook[analysis.name]
    date_column = _column_index(analysis, "A12 analysis_lia_date")
    assert analysis.column("A12").number_format == "yyyy-mm-dd"
    assert worksheet.cell(row=FIRST_DATA_ROW, column=date_column).number_format == "yyyy-mm-dd"


# ── 受控词表与固定枚举下拉 ──────────────────────────────────────────────────


def test_vocab_columns_have_dropdowns_and_resolvable_names(
    synthetic_profile, synthetic_layout
):
    """词表列有下拉，且 `Vocab_<id>` 命名区域可在工作簿里解析。"""
    workbook = _workbook(synthetic_profile)
    vocab = vocab_layout(synthetic_profile)
    assert vocab, "合成档案至少引用一个受控词表"

    checked = 0
    for sheet_spec in synthetic_layout.entry_sheets:
        worksheet = workbook[sheet_spec.name]
        for column in sheet_spec.columns:
            if not column.vocab_id:
                continue
            assert f"Vocab_{column.vocab_id}" in workbook.defined_names, column.vocab_id
            if column.value_kind is not ValueKind.CONTROLLED_VOCAB:
                continue  # 固定枚举列走系统命名区域 / 内联列表，另有测试
            validation = _validation_for(
                worksheet, _column_index(sheet_spec, column.header_key)
            )
            assert validation is not None, column.qualified_key
            assert validation.type == "list"
            assert validation.formula1 == f"=Vocab_{column.vocab_id}"
            checked += 1
    assert checked >= 2


def test_vocab_layout_lists_referenced_ids_and_terms(synthetic_profile):
    """`vocab_layout()` 覆盖被引用的词表，并带上 `vocab.py` 的词表项。"""
    layout = build_layout(synthetic_profile)
    referenced = {
        column.vocab_id
        for sheet in layout.sheets
        for column in sheet.columns
        if column.vocab_id
    }
    vocab = vocab_layout(synthetic_profile)
    assert referenced <= set(vocab)
    assert all(isinstance(items, list) for items in vocab.values())

    vocab_module = pytest.importorskip("data.metadata_profile.vocab")
    assert vocab["lia_ratio_name"] == list(vocab_module.terms("lia_ratio_name"))
    assert vocab["chemistry_method"] == list(vocab_module.terms("chemistry_method"))


def test_open_ended_vocab_allows_free_text_but_closed_rejects(
    synthetic_profile, synthetic_layout
):
    """开放词表的下拉不报错（允许自由文本），封闭词表严格拒绝。"""
    workbook = _workbook(synthetic_profile)
    sheet_spec = synthetic_layout.sheet_for_block(Block.LIA_RATIO)
    worksheet = workbook[sheet_spec.name]

    uncertainty_type = _validation_for(
        worksheet, _column_index(sheet_spec, "B6.3 lia_ratio_uncertainty_type")
    )
    assert uncertainty_type is not None
    assert uncertainty_type.formula1 == "=Vocab_lia_uncertainty_type"
    assert uncertainty_type.showErrorMessage is False, "开放词表必须允许自由文本"

    sigma = _validation_for(
        worksheet, _column_index(sheet_spec, "B6.4 lia_ratio_uncertainty_sigma")
    )
    assert sigma is not None and sigma.showErrorMessage is True


def test_fixed_enum_columns_have_dropdowns(synthetic_profile, synthetic_layout):
    """非词表枚举：SIGMA / RATIO_NAME / RATIO_SOURCE 各有下拉。"""
    workbook = _workbook(synthetic_profile)
    sheet_spec = synthetic_layout.sheet_for_block(Block.LIA_RATIO)
    worksheet = workbook[sheet_spec.name]

    sigma = _validation_for(
        worksheet, _column_index(sheet_spec, "B6.4 lia_ratio_uncertainty_sigma")
    )
    assert sigma is not None
    assert sigma.formula1 == '"' + ",".join(str(level) for level in SIGMA_LEVELS) + '"'
    assert sigma.showErrorMessage is True

    ratio_name = _validation_for(
        worksheet, _column_index(sheet_spec, "B6.1 lia_ratio_name")
    )
    assert ratio_name is not None
    assert ratio_name.formula1 == "=LIA_Ratio_Names"
    ratio_sheet = workbook[SHEET_VOCAB]
    ratio_column = get_column_letter(ratio_sheet.max_column - 1)
    assert workbook.defined_names["LIA_Ratio_Names"].attr_text == (
        f"'{SHEET_VOCAB}'!${ratio_column}$2:${ratio_column}${1 + len(LIA_RATIO_NAMES)}"
    )
    assert [
        ratio_sheet.cell(row=row, column=ratio_sheet.max_column - 1).value
        for row in range(2, 2 + len(LIA_RATIO_NAMES))
    ] == list(LIA_RATIO_NAMES)

    source = _validation_for(
        worksheet, _column_index(sheet_spec, "B6.7 lia_ratio_source")
    )
    assert source is not None
    assert source.formula1 == '"' + ",".join(RATIO_SOURCES) + '"'


def test_helper_columns_have_dropdowns(synthetic_profile, synthetic_layout):
    """辅助列：owner_sheet 用实体表命名区域，ratio_host 用档案宿主字段。"""
    workbook = _workbook(synthetic_profile)
    sheet_spec = synthetic_layout.sheet_for_block(Block.LIA_RATIO)
    worksheet = workbook[sheet_spec.name]

    owner_sheet = _validation_for(
        worksheet, _column_index(sheet_spec, OWNER_SHEET_COLUMN)
    )
    assert owner_sheet is not None
    assert owner_sheet.formula1 == f"={build_entry_workbook.OWNER_SHEET_LIST_NAME}"
    options_range = workbook.defined_names[build_entry_workbook.OWNER_SHEET_LIST_NAME]
    assert options_range.attr_text.startswith(f"'{SHEET_VOCAB}'!")
    vocab_sheet = workbook[SHEET_VOCAB]
    owner_options = [
        vocab_sheet.cell(row=row, column=vocab_sheet.max_column).value
        for row in range(2, 2 + len(owner_sheet_options(synthetic_layout)))
    ]
    assert owner_options == list(owner_sheet_options(synthetic_layout))
    # 明细行表也在 owner_sheet 下拉里（块表可以挂到具体的明细行上）
    row_sheet_names = {
        sheet.name
        for sheet in synthetic_layout.sheets
        if sheet.kind is SheetKind.GROUP
    }
    assert row_sheet_names and row_sheet_names <= set(owner_options)

    ratio_host = _validation_for(
        worksheet, _column_index(sheet_spec, RATIO_HOST_COLUMN)
    )
    assert ratio_host is not None
    assert ratio_host.showErrorMessage is True
    for host_oid in sheet_spec.host_oids:
        assert host_oid in ratio_host.formula1


def test_detail_row_sheets_have_group_dropdown(synthetic_profile, synthetic_layout):
    """明细行表的三个辅助列齐全，且 `group` 的下拉取值 == 布局给出的组根 OID。"""
    workbook = _workbook(synthetic_profile)
    checked = 0
    for sheet_spec in synthetic_layout.sheets:
        if sheet_spec.kind is not SheetKind.GROUP:
            continue
        assert sheet_spec.helper_columns == ("row_id", "parent_id", "group"), sheet_spec.name
        worksheet = workbook[sheet_spec.name]
        group_column = _column_index(sheet_spec, "group")
        validation = _validation_for(worksheet, group_column)
        assert validation is not None, sheet_spec.name
        assert validation.formula1 == '"' + ",".join(sheet_spec.host_oids) + '"'
        assert validation.showErrorMessage is True
        assert sheet_spec.host_oids, sheet_spec.name
        checked += 1
    assert checked >= 1, "合成档案应至少有一张明细行表"


def test_missing_vocab_module_still_builds_workbook(
    synthetic_profile, monkeypatch
):
    """`vocab.py` 缺失时回退为"无词表"：工作簿照常生成，命名区域仍在。"""
    monkeypatch.setattr(build_entry_workbook, "_vocab_module", None)
    monkeypatch.setattr(build_entry_workbook, "_vocab_module_checked", True)

    vocab = vocab_layout(synthetic_profile)
    assert vocab and all(items == [] for items in vocab.values())
    workbook = load_workbook(BytesIO(workbook_bytes(synthetic_profile)))
    for vocab_id in vocab:
        assert f"Vocab_{vocab_id}" in workbook.defined_names


# ── 必填条件格式 ────────────────────────────────────────────────────────────


def test_required_columns_are_highlighted_and_system_columns_are_not(
    synthetic_profile, synthetic_layout
):
    """必填且须录入者提供的列有红底规则；系统提供的列（B6.7）不得标红。"""
    workbook = _workbook(synthetic_profile)
    highlighted = 0
    for sheet_spec in synthetic_layout.entry_sheets:
        worksheet = workbook[sheet_spec.name]
        for column in sheet_spec.columns:
            rules = _conditional_rules_for(
                worksheet, _column_index(sheet_spec, column.header_key)
            )
            if column.is_required_from_provider:
                assert rules, f"{sheet_spec.name}/{column.qualified_key} 应有必填标记"
                highlighted += 1
            else:
                assert not rules, f"{sheet_spec.name}/{column.qualified_key} 不应标红"
    assert highlighted >= 3

    lia_spec = synthetic_layout.sheet_for_block(Block.LIA_RATIO)
    source = lia_spec.column("B6.7")
    assert source is not None and source.is_system_provided
    assert source.is_required_from_provider is False
    assert _conditional_rules_for(workbook[lia_spec.name], _column_index(lia_spec, source.header_key)) == []


def test_required_formula_only_fires_when_the_row_has_other_values(
    synthetic_profile, synthetic_layout
):
    """条件格式公式引用该列左右两侧的实际列字母。

    这里的字面公式是人核过的契约（Excel COM 复核一致）：17_LIA-Ratio 是
    ``owner_sheet / owner_id / ratio_host`` + ``B6.1..B6.7``，故 B6.1 落在 D 列；
    若块表的辅助列/列序变了，这条会失败，需重新核对公式而不是直接改期望值。
    """
    workbook = _workbook(synthetic_profile)
    lia_spec = synthetic_layout.sheet_for_block(Block.LIA_RATIO)
    name_column = _column_index(lia_spec, "B6.1 lia_ratio_name")
    assert get_column_letter(name_column) == "D"
    worksheet = workbook[lia_spec.name]
    rules = _conditional_rules_for(worksheet, name_column)
    assert len(rules) == 1
    assert rules[0].formula == [
        "AND(OR(COUNTA($A3:$C3)>0,COUNTA($E3:$J3)>0),ISBLANK(D3))"
    ]


def test_dropdown_and_highlight_ranges_cover_the_preset_rows(
    synthetic_profile, synthetic_layout
):
    """下拉与条件格式的作用域正好覆盖预设数据行（第 3 行到 202 行），列字母由布局派生。"""
    workbook = _workbook(synthetic_profile)
    sheet_spec = synthetic_layout.sheet_for_block(Block.LIA_RATIO)
    worksheet = workbook[sheet_spec.name]

    sigma_column = _column_index(sheet_spec, "B6.4 lia_ratio_uncertainty_sigma")
    sigma = _validation_for(worksheet, sigma_column)
    assert sigma is not None
    letter = get_column_letter(sigma_column)
    assert str(sigma.sqref) == f"{letter}{FIRST_DATA_ROW}:{letter}{LAST_DATA_ROW}"

    name_column = _column_index(sheet_spec, "B6.1 lia_ratio_name")
    name_letter = get_column_letter(name_column)
    rules = _conditional_rules_for(worksheet, name_column)
    ranges = [
        str(cell_range)
        for formatting in worksheet.conditional_formatting
        for cell_range in formatting.sqref.ranges
        if cell_range.min_col == name_column
    ]
    assert ranges == [f"{name_letter}{FIRST_DATA_ROW}:{name_letter}{LAST_DATA_ROW}"]
    assert rules


# ── 表头批注 ────────────────────────────────────────────────────────────────


def test_second_header_row_carries_field_comments(synthetic_profile, synthetic_layout):
    """第 2 行每个字段列都有批注，含定义/允许值/示例/来源，且框足够大。

    批注尺寸不在 xlsx 往返里保留，故这里查内存态工作簿（`_build_workbook`）。
    """
    workbook = _build_workbook(synthetic_profile)
    lia_spec = synthetic_layout.sheet_for_block(Block.LIA_RATIO)
    worksheet = workbook[lia_spec.name]
    for column in lia_spec.columns:
        cell = worksheet.cell(
            row=2, column=_column_index(lia_spec, column.header_key)
        )
        comment = cell.comment
        assert comment is not None, column.qualified_key
        assert column.definition in comment.text
        assert "来源:" in comment.text
        assert "中文标签:" in comment.text and "英文标签:" in comment.text
        assert comment.width >= 300 and comment.height >= 120
    for index in range(1, len(lia_spec.helper_columns) + 1):
        assert worksheet.cell(row=2, column=index).comment is None


# ── 0_SCHEMA / 0_VOCAB / 99_CHECK / README ─────────────────────────────────


def test_schema_rows_match_table_field_totals(synthetic_profile):
    """口径核对：`0_SCHEMA` 行数 == `sum(len(t.fields) for t in profile.tables)`。

    `Profile.fields` 本身就是"逐表逐字段（含内联副本）"的拼接，因此三个数必须相等；
    若哪天解析器改成"块字段每表只留一份"，这条会先失败，提醒两边一起改。
    """
    profile = synthetic_profile
    rows = schema_rows(profile)
    total = sum(len(table.fields) for table in profile.tables)
    assert len(rows) == total == len(profile.fields)
    assert len(rows) == _expected_schema_row_count(profile)


def test_schema_sheet_has_one_row_per_field_and_host(synthetic_profile):
    """`0_SCHEMA` 行数 == 字段按宿主展开后的条数，且表头固定。"""
    rows = schema_rows(synthetic_profile)
    assert len(rows) == len(synthetic_profile.fields)
    assert len(rows) == _expected_schema_row_count(synthetic_profile)
    assert all(len(row) == len(SCHEMA_HEADERS) for row in rows)

    workbook = _workbook(synthetic_profile)
    worksheet = workbook[SHEET_SCHEMA]
    assert [cell.value for cell in worksheet[1]] == list(SCHEMA_HEADERS)
    assert worksheet.max_row == len(rows) + 1
    assert worksheet.max_column == len(SCHEMA_HEADERS)
    assert worksheet.freeze_panes == "A2"
    last_column = get_column_letter(len(SCHEMA_HEADERS))
    assert worksheet.auto_filter.ref == f"A1:{last_column}{len(rows) + 1}"

    by_oid = {row[0]: row for row in rows}
    assert by_oid["B6.1"][4] == "analyses @ A14"
    assert by_oid["SI1"][5] == "mandatory"
    assert by_oid["SI1"][8] == "free_text"


def test_schema_expands_block_rows_per_host(synthetic_profile):
    """内联块字段按宿主各出一行，用"所属模块 @ 宿主"区分。"""
    rows = schema_rows(synthetic_profile)
    hosts = [row[4] for row in rows]
    assert hosts.count("analyses @ A7") == 2  # B4 的两个字段
    assert hosts.count("analyses @ A14") == 7  # B6 的七个子字段

    # 两个宿主各一行时总行数随之增加，仍等于独立算出的展开数
    from dataclasses import replace

    from data.metadata_profile.spec import TableSpec

    analysis = synthetic_profile.table("analyses")
    assert analysis is not None
    extra = replace(analysis, fields=analysis.fields + (analysis.field("B6.1"),))
    augmented = Profile(
        version=synthetic_profile.version,
        tables=(synthetic_profile.tables[0], extra),
        blocks=synthetic_profile.blocks,
    )
    assert len(schema_rows(augmented)) == _expected_schema_row_count(augmented)


def test_vocab_sheet_lists_columns_and_system_lists(synthetic_profile, synthetic_layout):
    """`0_VOCAB`：每张词表一列（含空词表），外加 ratio_name / owner_sheet 两列系统清单。"""
    workbook = _workbook(synthetic_profile)
    worksheet = workbook[SHEET_VOCAB]
    vocab = vocab_layout(synthetic_profile)
    assert worksheet.max_column == len(vocab) + 2
    assert worksheet.cell(row=1, column=1).value == "analysis_type 分析类型"
    assert worksheet.cell(row=1, column=len(vocab) + 1).value.startswith("ratio_name")
    assert worksheet.cell(row=1, column=len(vocab) + 2).value.startswith(OWNER_SHEET_COLUMN)
    assert worksheet.freeze_panes == "A2"
    assert len(workbook.defined_names) == len(vocab) + 2


def test_check_sheet_header_comes_from_io_xlsx(synthetic_profile):
    """`99_CHECK` 表头与 `io_xlsx.CHECK_HEADERS` 一致（只有一层表头）。"""
    workbook = _workbook(synthetic_profile)
    worksheet = workbook[SHEET_CHECK]
    assert [cell.value for cell in worksheet[1]] == list(CHECK_HEADERS)
    # 先读 max_row 再访问 A2 —— openpyxl 的 cell() 会创建单元格并撑大 max_row
    assert worksheet.max_row == 1
    assert worksheet["A2"].value is None


def test_readme_explains_the_five_required_topics(synthetic_profile):
    """README 覆盖：生成物勿改 / 实体链路 / 可重复块 / ratio_host / 必填标记与校验。"""
    workbook = _workbook(synthetic_profile)
    worksheet = workbook[SHEET_README]
    text = "\n".join(
        str(worksheet.cell(row=row, column=1).value or "")
        for row in range(1, worksheet.max_row + 1)
    )
    for keyword in (
        "reference/metadata",
        "0_SCHEMA",
        "0_VOCAB",
        "99_CHECK",
        "勿手改",
        "Site",
        "Assemblage",
        "Object",
        "Sample",
        "Analysis",
        "owner_sheet",
        "owner_id",
        "ratio_host",
        "A14",
        "★",
        "校验",
    ):
        assert keyword in text, keyword


# ── 落盘与命令行 ────────────────────────────────────────────────────────────


def test_build_workbook_writes_file_and_creates_parent(tmp_path, synthetic_profile):
    """`build_workbook()` 落盘并返回路径，父目录自动创建。"""
    target = tmp_path / "nested" / "dir" / "entry.xlsx"
    written = build_workbook(synthetic_profile, target)
    assert written == target
    assert target.exists()
    assert load_workbook(target).sheetnames


def test_main_refuses_to_overwrite_without_force(
    tmp_path, synthetic_profile, monkeypatch, capsys
):
    """目标已存在且未给 `--force` 时拒绝覆盖，并 `return 1`。"""
    monkeypatch.setattr(build_entry_workbook, "_load_profile", lambda root: synthetic_profile)
    target = tmp_path / "filled.xlsx"
    target.write_bytes(b"user data")

    assert main(["--output", str(target)]) == 1
    assert target.read_bytes() == b"user data"
    assert "已存在" in capsys.readouterr().out

    assert main(["--output", str(target), "--force"]) == 0
    assert load_workbook(target).sheetnames


def test_main_default_output_name_uses_profile_version(
    tmp_path, synthetic_profile, monkeypatch
):
    """未给 `--output` 时输出 `entry_template_TerraLID_v<版本>.xlsx`。"""
    monkeypatch.setattr(build_entry_workbook, "_load_profile", lambda root: synthetic_profile)
    monkeypatch.setattr(build_entry_workbook, "DEFAULT_OUTPUT_DIR", tmp_path / "数据")

    assert main(["--force"]) == 0
    expected = tmp_path / "数据" / f"entry_template_TerraLID_v{synthetic_profile.version}.xlsx"
    assert expected.exists()


def test_main_emits_registry_json_and_prints_stats(
    tmp_path, synthetic_profile, synthetic_layout, monkeypatch, capsys
):
    """`--emit-json` 输出注册表快照，stdout 打印输出路径与统计。"""
    monkeypatch.setattr(build_entry_workbook, "_load_profile", lambda root: synthetic_profile)
    target = tmp_path / "entry.xlsx"
    snapshot_path = tmp_path / "registry.json"

    code = main(
        [
            "--output",
            str(target),
            "--emit-json",
            str(snapshot_path),
            "--force",
        ]
    )
    assert code == 0
    snapshot = json.loads(snapshot_path.read_text(encoding="utf-8"))
    assert snapshot["profile_version"] == synthetic_profile.version
    assert snapshot["sheet_count"] == len(synthetic_layout.sheets)
    assert snapshot["schema_row_count"] == len(schema_rows(synthetic_profile))
    assert [sheet["name"] for sheet in snapshot["sheets"]] == list(
        synthetic_layout.sheet_names
    )
    assert set(snapshot["vocabularies"]) == set(vocab_layout(synthetic_profile))

    out = capsys.readouterr().out
    assert str(target) in out
    assert "工作表" in out and "受控词表" in out


def test_main_reports_failure_when_profile_cannot_be_loaded(
    tmp_path, monkeypatch, capsys
):
    """解析器不可用时 `main()` 打印错误并 `return 1`（不抛栈）。"""
    def _boom(root: Path) -> Profile:
        raise RuntimeError("no parser")

    monkeypatch.setattr(build_entry_workbook, "_load_profile", _boom)
    assert main(["--output", str(tmp_path / "x.xlsx")]) == 1
    assert "加载注册表失败" in capsys.readouterr().out


# ── 真实档案端到端（依赖 task-1 的 load_profile） ──────────────────────────


@pytest.fixture(scope="module")
def real_profile() -> Profile:
    """真实 TerraLID 档案的注册表（parser 与 vocab/labels 都已落地）。"""
    try:
        from data.metadata_profile.profile import load_profile
    except ImportError as err:  # pragma: no cover - 解析器尚未落地时跳过
        pytest.skip(f"档案解析器不可用: {err}")
    return load_profile()


def test_real_profile_workbook_matches_layout(real_profile):
    """真实档案生成的工作簿与布局逐张一致（表数/字段数不写死，只与布局和注册表对齐）。"""
    layout = build_layout(real_profile)
    workbook = _workbook(real_profile)
    assert workbook.sheetnames == list(layout.sheet_names)
    assert len(layout.sheets) == len(workbook.sheetnames)

    rows = schema_rows(real_profile)
    assert len(rows) == sum(len(table.fields) for table in real_profile.tables)
    assert len(rows) == _expected_schema_row_count(real_profile)
    assert workbook[SHEET_SCHEMA].max_row == len(rows) + 1


def test_real_profile_primary_key_columns_do_not_promise_auto_fill(real_profile):
    """5 个 TerraLID ID 主键列：批注说明"工具不代为生成"，且仍不标红。

    档案把它们的 `Provided by` 记为 `TerraLID system`，但工具并不生成这些 ID —— 批注不能
    承诺"可留空"，否则用户留空后块表/明细行表会无值可挂。
    """
    layout = build_layout(real_profile)
    workbook = _build_workbook(real_profile)  # 批注只在内存态可完整读取
    for sheet_name in ("1_Site", "2_Assemblage", "3_Object", "4_Sample", "5_Analysis"):
        sheet_spec = layout.sheet(sheet_name)
        assert sheet_spec is not None
        key_column = sheet_spec.columns[0]
        assert key_column.key.startswith("terralid_")
        # 主键是第一个**字段**列；有父外键的表前面还有 helper 列
        key_index = len(sheet_spec.helper_columns) + 1
        assert sheet_spec.all_header_keys[key_index - 1] == key_column.header_key
        comment = workbook[sheet_name].cell(row=2, column=key_index).comment
        assert comment is not None, sheet_name
        assert "TerraLID 数据库分配" in comment.text, sheet_name
        assert "本工具不代为生成" in comment.text, sheet_name
        assert "自动补齐" not in comment.text, sheet_name
        assert "可留空" not in comment.text, sheet_name
        # 系统提供 → 不得标红（is_system_provided 语义不变）
        assert key_column.is_required_from_provider is False
        assert _conditional_rules_for(workbook[sheet_name], key_index) == [], sheet_name

    # 真正由工具补齐的列仍保留"可留空"口径
    lia_spec = layout.sheet_for_block(Block.LIA_RATIO)
    source_comment = workbook[lia_spec.name].cell(
        row=2, column=_column_index(lia_spec, "B6.7 lia_ratio_source")
    ).comment
    assert source_comment is not None and "自动补齐" in source_comment.text


def test_real_profile_readme_does_not_claim_ids_are_auto_filled(real_profile):
    """README 的主键说明与批注同一口径，且不再把 ID 列说成"自动补齐/可留空"。"""
    workbook = _workbook(real_profile)
    worksheet = workbook[SHEET_README]
    text = "\n".join(
        str(worksheet.cell(row=row, column=1).value or "")
        for row in range(1, worksheet.max_row + 1)
    )
    assert "本工具不代为生成" in text
    for oid in ("SI0", "AS0", "O0", "S0", "A0"):
        assert oid in text, oid
    assert "自动补齐" in text  # 用于 B6.7 等真正补齐的列
    assert "工具【不会】代填" in text
    assert "主键（TerraLID ID）列" in text
    assert "**" not in text, "README 单元格是纯文本，不能出现 Markdown 星号"


def _group_cells(sheet_spec, group_oid: str) -> list[str]:
    """分组在本表的各列字母，按列序（本行固定第 3 行）。"""
    prefix = f"{group_oid}."
    offset = len(sheet_spec.helper_columns)
    return [
        f"${get_column_letter(index)}{FIRST_DATA_ROW}"
        for index, candidate in enumerate(sheet_spec.columns, start=offset + 1)
        if candidate.oid == group_oid or candidate.oid.startswith(prefix)
    ]


def _nearest_optional_group(column) -> str | None:
    """该列最近的可选分组（真相源 = `layout.ColumnSpec.optional_ancestors`，自近及远）。

    跳过可复用块的块根（`B1`…`B6`）：块表一行本身就是"用了这一组"的声明，块根不该成为
    条件 —— 否则 `B6.1` 的因子会自我指涉（`COUNTA(B6.1..B6.7)>0`），空比值行永不标红。
    """
    block_oids = {column.block.value} if column.block is not None else set()
    return next(
        (ancestor for ancestor in column.optional_ancestors if ancestor not in block_oids),
        None,
    )


def _expected_group_factor(sheet_spec, column) -> str:
    """独立算出"分组已启用"因子：该列最近的可选分组在本行的各列。"""
    group_oid = _nearest_optional_group(column)
    assert group_oid is not None, column.qualified_key
    cells = _group_cells(sheet_spec, group_oid)
    assert cells, group_oid
    return f"COUNTA({','.join(cells)})>0"


def _formula_for(workbook, sheet_spec, oid: str) -> list:
    """取某列的条件格式公式文本。"""
    index = _column_index(sheet_spec, sheet_spec.column(oid).header_key)
    return [
        formula
        for rule in _conditional_rules_for(workbook[sheet_spec.name], index)
        for formula in rule.formula
    ]


def test_real_profile_optional_groups_make_required_formatting_conditional(real_profile):
    """可选分组内的必填列多一条"该组已启用"因子；组外必填列保持原公式。"""
    layout = build_layout(real_profile)
    workbook = _workbook(real_profile)

    conditional: list[tuple[str, str]] = []
    unconditional: list[tuple[str, str]] = []
    for sheet_spec in layout.entry_sheets:
        for column in sheet_spec.columns:
            if not column.is_required_from_provider:
                continue
            target = conditional if _nearest_optional_group(column) else unconditional
            target.append((sheet_spec.name, column.oid))
    assert len(conditional) >= 20, "档案里有大量可选分组内的必填列"
    assert len(unconditional) >= 5

    for sheet_name, oid in conditional:
        sheet_spec = layout.sheet(sheet_name)
        formulas = _formula_for(workbook, sheet_spec, oid)
        assert len(formulas) == 1, f"{sheet_name}/{oid}"
        factor = _expected_group_factor(sheet_spec, sheet_spec.column(oid))
        assert formulas[0].startswith(f"AND({factor},"), f"{sheet_name}/{oid}: {formulas[0]}"

    for sheet_name, oid in unconditional:
        sheet_spec = layout.sheet(sheet_name)
        column = sheet_spec.column(oid)
        formulas = _formula_for(workbook, sheet_spec, oid)
        assert len(formulas) == 1, f"{sheet_name}/{oid}"
        assert not formulas[0].startswith("AND(COUNTA($"), (
            f"{sheet_name}/{oid} 不应带分组因子: {formulas[0]}"
        )


def test_real_profile_group_factor_covers_only_its_own_group(real_profile):
    """抽查 `1_Site` 的 SI5.1.1 / SI5.2.1：因子只含本组各列，不含同级的另一组。"""
    layout = build_layout(real_profile)
    workbook = _workbook(real_profile)
    site = layout.sheet("1_Site")

    point_factor = _expected_group_factor(site, site.column("SI5.1.1"))
    bbox_factor = _expected_group_factor(site, site.column("SI5.2.1"))
    assert point_factor == f"COUNTA({','.join(_group_cells(site, 'SI5.1'))})>0"
    assert bbox_factor == f"COUNTA({','.join(_group_cells(site, 'SI5.2'))})>0"

    point_formula = _formula_for(workbook, site, "SI5.1.1")[0]
    bbox_formula = _formula_for(workbook, site, "SI5.2.1")[0]
    assert point_formula.startswith(f"AND({point_factor},"), point_formula
    assert bbox_formula.startswith(f"AND({bbox_factor},"), bbox_formula
    # 点位组的因子不得把四至组（SI5.2）算作"已启用"
    for letter in [cell[1] for cell in _group_cells(site, "SI5.2")]:
        assert f"${letter}3" not in point_factor, point_factor
    assert point_factor != bbox_factor


def test_real_profile_group_free_columns_keep_unconditional_formula(real_profile):
    """lead 点名的无条件必填列不得带分组因子。"""
    layout = build_layout(real_profile)
    workbook = _workbook(real_profile)
    for sheet_name, oid in (
        ("1_Site", "SI1"),
        ("3_Object", "O3"),
        ("3_Object", "O12"),
        ("4_Sample", "S5"),
        ("4_Sample", "S8"),
        ("5_Analysis", "A2"),
        ("5_Analysis", "A6.1"),
        ("17_LIA-Ratio", "B6.1"),
        ("17_LIA-Ratio", "B6.2"),
    ):
        sheet_spec = layout.sheet(sheet_name)
        formulas = _formula_for(workbook, sheet_spec, oid)
        assert len(formulas) == 1, f"{sheet_name}/{oid}"
        assert formulas[0].startswith("AND(OR("), f"{sheet_name}/{oid}: {formulas[0]}"


def test_readme_minimal_entry_section_is_registry_derived(real_profile):
    """README 第一节「零、最小填写路径」：列名由注册表现算，且排在其它小节之前。"""
    layout = build_layout(real_profile)
    workbook = _workbook(real_profile)
    worksheet = workbook[SHEET_README]
    lines = [
        str(worksheet.cell(row=row, column=1).value or "")
        for row in range(1, worksheet.max_row + 1)
    ]
    section_zero = next(i for i, line in enumerate(lines) if line.startswith("零、最小填写路径"))
    section_one = next(i for i, line in enumerate(lines) if line.startswith("一、哪些表不能手改"))
    assert section_zero < section_one
    block = "\n".join(lines[section_zero:section_one])

    # 绝对最小：分析表的必填列 + 比值块表的辅助列与必填列（全部由注册表推出）
    assert "5_Analysis" in block and "A2" in block and "A6.1" in block
    assert "17_LIA-Ratio" in block
    for helper in ("owner_sheet", "owner_id", "ratio_host"):
        assert helper in block
    assert "B6.1" in block and "B6.2" in block
    # 主键提示与"其余 N 张录入表"
    for oid in ("SI0", "AS0", "O0", "S0", "A0"):
        assert oid in block
    assert f"其余 {len(layout.entry_sheets) - 2} 张录入表" in block
    # 推荐最小：5 格无条件必填
    for oid in ("SI1", "O3", "O12", "S5", "S8"):
        assert oid in block
    # 条件必填规则与收尾提醒
    assert "可选分组" in block and "0–1 / 0–n" in block
    assert "99_CHECK" in block
    assert "**" not in block


def test_real_profile_has_both_ratio_hosts_and_all_helper_dropdowns(real_profile):
    """真实档案里 ratio_host 同时给出 A9.3 与 A14，owner_sheet 给出全部实体表。"""
    layout = build_layout(real_profile)
    workbook = _workbook(real_profile)
    lia_spec = layout.sheet_for_block(Block.LIA_RATIO)
    assert lia_spec is not None
    worksheet = workbook[lia_spec.name]

    ratio_host = _validation_for(worksheet, _column_index(lia_spec, RATIO_HOST_COLUMN))
    assert ratio_host is not None
    assert "A9.3" in ratio_host.formula1 and "A14" in ratio_host.formula1

    owner_sheet = _validation_for(worksheet, _column_index(lia_spec, OWNER_SHEET_COLUMN))
    assert owner_sheet is not None
    assert owner_sheet.formula1 == f"={build_entry_workbook.OWNER_SHEET_LIST_NAME}"
    vocab_sheet = workbook[SHEET_VOCAB]
    options = [
        vocab_sheet.cell(row=row, column=vocab_sheet.max_column).value
        for row in range(2, 2 + len(owner_sheet_options(layout)))
    ]
    assert options == list(owner_sheet_options(layout))


def test_real_profile_covers_remaining_fixed_enums(real_profile):
    """只有真档案才有的固定枚举（GEOL_TYPE / UNIT_AGE / YES_NO_UNCLEAR）也有下拉。"""
    layout = build_layout(real_profile)
    workbook = _workbook(real_profile)
    expected = {
        ValueKind.GEOL_TYPE: '"geological,archaeological"',
        ValueKind.UNIT_AGE: '"a,Ma"',
        ValueKind.YES_NO_UNCLEAR: '"yes,no,unclear"',
    }
    found: set[ValueKind] = set()
    for sheet_spec in layout.entry_sheets:
        worksheet = workbook[sheet_spec.name]
        for column in sheet_spec.columns:
            formula = expected.get(column.value_kind)
            if formula is None:
                continue
            validation = _validation_for(
                worksheet, _column_index(sheet_spec, column.header_key)
            )
            assert validation is not None, column.qualified_key
            assert validation.formula1 == formula, column.qualified_key
            found.add(column.value_kind)
    assert found == set(expected)


def test_real_profile_cli_end_to_end(tmp_path):
    """`main()` 跑真实档案：写临时路径成功、打印统计、`--emit-json` 可读。"""
    target = tmp_path / "entry_template.xlsx"
    snapshot_path = tmp_path / "registry.json"
    code = main(
        [
            "--output",
            str(target),
            "--emit-json",
            str(snapshot_path),
            "--force",
        ]
    )
    assert code == 0
    workbook = load_workbook(target)
    assert workbook.sheetnames[0] == SHEET_README
    assert workbook.sheetnames[-1] == SHEET_CHECK
    snapshot = json.loads(snapshot_path.read_text(encoding="utf-8"))
    assert snapshot["sheet_count"] == len(workbook.sheetnames)
    assert snapshot["schema_row_count"] == workbook[SHEET_SCHEMA].max_row - 1
    assert snapshot["vocabularies"]
    assert any(
        sheet["kind"] == SheetKind.BLOCK.value for sheet in snapshot["sheets"]
    )
